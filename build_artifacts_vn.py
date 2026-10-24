from __future__ import annotations

import json
import math
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Inches as PInches, Pt as PPt


BASE = Path(r"C:\Users\dangv\Downloads")
REPO = Path(r"C:\Users\dangv\Desktop\SIC\SIC-Smart-Attendance")
OUT = Path.cwd() / "deliverables"
ASSETS = OUT / "report_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)


def get_font(size: int, bold: bool = False):
    filename = "arialbd.ttf" if bold else "arial.ttf"
    path = Path(rf"C:\Windows\Fonts\{filename}")
    if not path.exists():
        path = Path(r"C:\Windows\Fonts\segoeui.ttf")
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def make_diagram(name: str, title: str, labels: list[str]) -> Path:
    image = Image.new("RGB", (1500, 800), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), title, fill="#102A43", font=get_font(38, True))
    box_w, box_h = 355, 115
    positions = []
    for index, label in enumerate(labels):
        row, col = divmod(index, 3)
        x, y = 70 + col * 480, 150 + row * 220
        positions.append((x, y))
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=18, fill="#D9EAF7", outline="#1F4E79", width=4)
        lines = label.split("\n")
        for line_no, line in enumerate(lines):
            bbox = draw.textbbox((0, 0), line, font=get_font(23, True))
            tx = x + (box_w - bbox[2]) / 2
            ty = y + 32 + line_no * 30
            draw.text((tx, ty), line, fill="#102A43", font=get_font(23, True))
    for index in range(len(positions) - 1):
        x1, y1 = positions[index]
        x2, y2 = positions[index + 1]
        start = (x1 + box_w, y1 + box_h // 2)
        end = (x2, y2 + box_h // 2)
        if x2 < x1:
            start, end = (x1, y1 + box_h // 2), (x2 + box_w, y2 + box_h // 2)
        draw.line([start, end], fill="#2F80ED", width=5)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        tip = end
        left = (tip[0] - 22 * math.cos(angle - 0.5), tip[1] - 22 * math.sin(angle - 0.5))
        right = (tip[0] - 22 * math.cos(angle + 0.5), tip[1] - 22 * math.sin(angle + 0.5))
        draw.polygon([tip, left, right], fill="#2F80ED")
    path = ASSETS / f"{name}.png"
    image.save(path)
    return path


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    properties.append(node)


def margins(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), "100")
        item.set(qn("w:type"), "dxa")
        node.append(item)
    properties.append(node)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        cell = result.rows[0].cells[i]
        cell.text = value
        shade(cell, "1F4E79")
        margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cells[i])
            if row_index % 2:
                shade(cells[i], "EAF2F8")


def h(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def p(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.15


def list_items(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="Normal")
        # Use the bullet definition already supplied by the Samsung template
        # instead of inserting a literal bullet character.
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "2")
        num_pr.append(ilvl)
        num_pr.append(num_id)
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(3)


def figure(doc: Document, path: Path, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.2))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True


def configure(doc: Document) -> None:
    section = doc.sections[0]
    # Match the supplied Samsung report template: A4 portrait, 1 inch side
    # margins, 1.18 inch top margin, and the template's footer spacing.
    section.page_width = Emu(7560310)
    section.page_height = Emu(10692130)
    section.top_margin = Emu(1080135)
    section.bottom_margin = Emu(914400)
    section.left_margin = Emu(914400)
    section.right_margin = Emu(914400)
    section.header_distance = Emu(540385)
    section.footer_distance = Emu(144145)
    styles = doc.styles
    styles["Normal"].font.name = "SamsungOne-400"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Normal"].paragraph_format.space_after = Pt(9)
    for name, size in (("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 11)):
        styles[name].font.name = "SamsungOne-700"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.keep_with_next = True
    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        doc.settings.element.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build_report() -> Path:
    metric_path = BASE / "metrics.json"
    metrics = json.loads(metric_path.read_text(encoding="utf-8")) if metric_path.exists() else {}
    precision = metrics.get("precision", 0.8253197803081475)
    recall = metrics.get("recall", 0.5617649280983202)
    map50 = metrics.get("map50", 0.6437364126984133)
    map5095 = metrics.get("map50_95", 0.33972672829922285)
    # Latest ArcFace notebook metrics: face-attendance-arcface-recognition (1).ipynb
    arcface_val_top1 = 0.8895875519556645
    arcface_val_top5 = 0.9208142385164659
    arcface_gallery_top1 = 0.8801023126931685
    arcface_gallery_top5 = 0.9140999680272834
    arcface_true_cosine_mean = 0.7741613417548111
    arcface_margin_mean = 0.3712819725430674
    arcface_params = 27910327
    arcface_epochs = 20
    arcface_checkpoint = "embedder_best.pt"
    # Separate diagnostic on the model currently integrated in the system.
    system_top1 = 0.975
    system_top5 = 1.0
    system_mean_cosine = 0.7936312291771174
    system_min_cosine = 0.05940623581409454
    system_mean_margin = 0.33510992834344505
    architecture = make_diagram("architecture", "Kiến trúc SPAS", ["Webcam / camera", "Frontend", "Backend", "AI service", "PostgreSQL", "Evidence"])
    pipeline = make_diagram("pipeline", "Pipeline AI", ["Ảnh đầu vào", "YOLO11n", "Crop + chuẩn hóa", "Embedding 512D", "Cosine similarity", "Identity / UNKNOWN"])
    attendance = make_diagram("attendance", "Luồng điểm danh", ["GV chọn lớp", "Nạp roster", "Quét nhiều frame", "Lọc UNKNOWN", "Lưu attendance", "GV review"])

    # Start from the supplied Samsung report file so its theme, fonts,
    # numbering definitions, footer, and page-number fields are preserved.
    doc = Document(BASE / "SIC_AI_Capstone Project_Final Report.docx")
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    configure(doc)
    for _ in range(3):
        doc.add_paragraph()
    ai_course = doc.add_paragraph()
    ai_course_run = ai_course.add_run("AI Course")
    ai_course_run.font.name = "SamsungOne-700"
    ai_course_run.font.size = Pt(20)
    cover_title = doc.add_paragraph()
    cover_title_run = cover_title.add_run("Capstone Project\nFinal Report")
    cover_title_run.font.name = "Samsung Sharp Sans"
    cover_title_run.font.size = Pt(36)
    cover_title_run.font.color.rgb = RGBColor(25, 61, 176)
    project_title = doc.add_paragraph()
    project_title_run = project_title.add_run("Smart Passive Attendance System")
    project_title_run.font.name = "SamsungOne-700"
    project_title_run.font.size = Pt(20)
    project_title_run.font.color.rgb = RGBColor(0, 0, 0)
    project_title.paragraph_format.space_after = Pt(14)
    cover_note = doc.add_paragraph()
    cover_note_run = cover_note.add_run("For students (instructor review required)")
    cover_note_run.font.name = "SamsungOne-700"
    cover_note_run.font.size = Pt(16)
    cover_info = doc.add_paragraph()
    cover_info_run = cover_info.add_run("Nhóm: SIC AI Team\nTháng 09 năm 2026")
    cover_info_run.font.name = "SamsungOne-400"
    cover_info_run.font.size = Pt(11)
    for _ in range(5):
        doc.add_paragraph()
    for line in [
        "ⓒ2023 SAMSUNG. All rights reserved.",
        "Samsung Electronics Corporate Citizenship Office holds the copyright of this document.",
        "This document is a literary property protected by copyright law so reprint and reproduction without permission are prohibited.",
        "To use this document other than the curriculum of Samsung Innovation Campus, you must receive written consent from copyright holder.",
    ]:
        disclaimer = doc.add_paragraph(line)
        for run in disclaimer.runs:
            run.font.name = "SamsungOne-400"
            run.font.size = Pt(7)
        disclaimer.paragraph_format.space_after = Pt(0)
    doc.add_page_break()

    h(doc, "Tóm tắt")
    p(doc, "Smart Passive Attendance System (SPAS) là hệ thống hỗ trợ điểm danh trong lớp học bằng camera. Hệ thống tìm khuôn mặt, tạo vector đặc trưng, so sánh với danh sách sinh viên đã đăng ký và lưu kết quả cùng ảnh minh chứng.")
    p(doc, "SPAS tách AI thành hai bước. YOLO11n tìm vị trí khuôn mặt. FaceNet/ArcFace tạo vector 512 chiều. Cosine similarity so sánh vector mới với gallery của lớp. Khi điểm giống nhau thấp hoặc không rõ, hệ thống trả UNKNOWN thay vì tự gán sai người.")
    p(doc, "Báo cáo tập trung vào dữ liệu, preprocessing, model, metric và tích hợp. Metric YOLO11n và ArcFace được lấy từ notebook latest đã train. Ngoài ra, report có một phép đo nội bộ trên model đang cài trong hệ thống; hai nhóm số liệu được tách riêng để tránh nhầm lẫn.")
    table(doc, ["Nội dung", "Tóm tắt"], [["Bài toán", "Giảm thời gian điểm danh và có ảnh minh chứng."], ["Model 1", "YOLO11n phát hiện khuôn mặt."], ["Model 2", "FaceNet/ArcFace tạo embedding 512 chiều."], ["So khớp", "Cosine similarity với gallery theo roster."], ["Phạm vi", "MVP cho lớp học, có teacher review."]])
    doc.add_page_break()

    h(doc, "Mục lục")
    toc_entries = [
        ("Chương 1. Đặt vấn đề và tổng quan hệ thống", 4),
        ("Chương 2. Cơ sở lý thuyết", 6),
        ("Chương 3. Phân tích và thiết kế hệ thống", 9),
        ("Chương 4. Huấn luyện mô hình và tích hợp hệ thống", 15),
        ("Chương 5. Kết quả, kiểm thử và đánh giá hệ thống", 23),
        ("Kết luận và hướng phát triển", 29),
        ("Tài liệu tham khảo", 30),
        ("Phụ lục", 31),
    ]
    for entry, page_number in toc_entries:
        toc_line = doc.add_paragraph()
        toc_line.paragraph_format.left_indent = Inches(0.30)
        toc_line.paragraph_format.first_line_indent = Inches(-0.30)
        toc_line.paragraph_format.space_after = Pt(5)
        toc_line.paragraph_format.line_spacing = 1.15
        toc_line.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        toc_run = toc_line.add_run(f"{entry}\t{page_number}")
        toc_run.font.name = "SamsungOne-400"
        toc_run.font.size = Pt(11)
    doc.add_page_break()

    h(doc, "Đặt vấn đề")
    h(doc, "1. Bối cảnh", 2)
    p(doc, "Điểm danh thủ công mất thời gian và có thể xảy ra nhập nhầm. Khi có tranh chấp, giáo viên thường không có ảnh hoặc thông tin đủ rõ để kiểm tra lại. Camera trong lớp có thể cung cấp dữ liệu đầu vào cho một quy trình tự động hơn.")
    p(doc, "Nhóm chọn bài toán điểm danh khuôn mặt vì đây là bài toán AI dễ trình bày theo từng bước: dữ liệu, phát hiện, nhận diện, đánh giá và tích hợp. AI tạo kết quả ban đầu; giáo viên vẫn là người kiểm tra trường hợp bất thường.")
    h(doc, "2. Mục tiêu", 2)
    list_items(doc, ["Xây dựng pipeline phát hiện và nhận diện khuôn mặt.", "Huấn luyện hoặc fine-tune model trên dữ liệu phù hợp.", "Đo metric trên dữ liệu validation và dữ liệu hệ thống.", "Tích hợp model vào AI service và luồng điểm danh.", "Lưu kết quả cùng ảnh evidence để giáo viên kiểm tra."])
    h(doc, "3. Phạm vi", 2)
    p(doc, "MVP hỗ trợ đăng nhập, phân quyền, đăng ký khuôn mặt, quản lý lớp, camera/webcam, nhận diện người trong roster, UNKNOWN_PERSON, evidence và chỉnh sửa kết quả. Nhận diện hành vi, liveness chuyên sâu và triển khai nhiều camera được để ở hướng phát triển.")
    doc.add_page_break()

    h(doc, "Đặt vấn đề. Cách tiếp cận của nhóm")
    h(doc, "4. Vì sao cần AI hỗ trợ", 2)
    p(doc, "Mục tiêu của nhóm không phải thay thế giáo viên. Mục tiêu là tự động hóa phần quan sát lặp lại: tìm mặt trong ảnh, đề xuất identity và lưu evidence. Giáo viên vẫn quyết định với các trường hợp mờ, thiếu sáng hoặc không chắc chắn.")
    h(doc, "5. Câu hỏi AI của dự án", 2)
    table(doc, ["Câu hỏi", "Cách trả lời trong dự án"], [["Ảnh có bao nhiêu khuôn mặt?", "YOLO11n trả bounding box."], ["Khuôn mặt thuộc ai?", "Recognizer tạo embedding và so với gallery."], ["Khi nào không nên đoán?", "Dùng threshold và ambiguity margin để trả UNKNOWN."], ["Làm sao kiểm tra kết quả?", "Dùng metric notebook, test nội bộ và teacher review."]])
    h(doc, "6. Kết quả mong muốn", 2)
    p(doc, "Một phiên điểm danh cần tạo ra kết quả có thể giải thích: ảnh đầu vào, vị trí mặt, identity được đề xuất, điểm similarity, thời gian và ảnh evidence. Các trường hợp không chắc chắn phải được đánh dấu rõ thay vì âm thầm ghi nhận.")
    doc.add_page_break()

    h(doc, "Chương 1. Cơ sở lý thuyết")
    h(doc, "1.1. Computer Vision", 2)
    p(doc, "Computer Vision giúp máy tính hiểu hình ảnh hoặc video. Trong SPAS, ảnh được chuyển thành ba loại thông tin: vị trí khuôn mặt, vector đặc trưng và kết quả điểm danh.")
    h(doc, "1.2. Face Detection", 2)
    p(doc, "Face Detection là bước tìm bounding box của từng khuôn mặt. Model trả về vị trí và confidence. Bước này chạy trước nhận diện vì ảnh lớp học có thể có nhiều người ở nhiều vị trí.")
    h(doc, "1.3. CNN và transfer learning", 2)
    p(doc, "CNN học đặc trưng không gian trong ảnh. Layer đầu học cạnh và vùng sáng tối. Layer sâu học mẫu phức tạp hơn. Transfer learning bắt đầu từ weight đã học trên dữ liệu lớn, sau đó điều chỉnh cho bài toán mới. Cách này tiết kiệm dữ liệu và thời gian train.")
    h(doc, "1.4. Face embedding", 2)
    p(doc, "Embedding là vector số mô tả khuôn mặt. Ảnh cùng một người nên tạo vector gần nhau. Ảnh người khác nên tạo vector xa hơn. Hệ thống lưu các vector đăng ký trong gallery.")
    figure(doc, pipeline, "Hình 1.1. Pipeline AI từ ảnh đến identity")
    doc.add_page_break()

    h(doc, "Chương 1. Cơ sở lý thuyết tiếp theo")
    h(doc, "1.5. Cosine similarity", 2)
    p(doc, "Cosine similarity đo mức giống nhau về hướng giữa hai vector. Sau khi chuẩn hóa, hệ thống dùng tích vô hướng để tính điểm. Điểm cao nghĩa là hai vector giống nhau hơn. Threshold cần được kiểm tra trên dữ liệu thật, không nên chọn chỉ bằng cảm tính.")
    h(doc, "1.6. ArcFace", 2)
    p(doc, "ArcFace thêm angular margin vào lúc train. Margin làm vector cùng identity gom thành cụm chặt hơn và vector khác identity tách xa hơn. Khi deploy, hệ thống dùng embedding và gallery; classification head chủ yếu phục vụ training.")
    h(doc, "1.7. Metric", 2)
    table(doc, ["Metric", "Cách hiểu đơn giản"], [["Precision", "Trong các mặt model đoán, bao nhiêu mặt đúng."], ["Recall", "Trong các mặt thật, model tìm được bao nhiêu mặt."], ["mAP50", "Chất lượng bounding box ở IoU 0.50."], ["mAP50-95", "Chất lượng box ở nhiều ngưỡng chặt hơn."], ["Top-1", "Identity đứng đầu có đúng không."], ["Top-5", "Identity đúng có nằm trong 5 kết quả đầu không."]])
    h(doc, "1.8. Quyền riêng tư", 2)
    p(doc, "Dữ liệu khuôn mặt là dữ liệu nhạy cảm. Cần có consent, giới hạn quyền xem, bảo vệ ảnh và embedding, không ghi raw vector vào log, quy định thời gian lưu trữ và duy trì quyền review của giáo viên.")
    doc.add_page_break()

    h(doc, "Chương 1. Ví dụ minh họa dễ hiểu")
    h(doc, "1.9. Một ảnh đi qua hệ thống như thế nào", 2)
    list_items(doc, ["Camera gửi một frame có nhiều sinh viên.", "YOLO11n vẽ một box cho mỗi khuôn mặt tìm được.", "Mỗi box được crop và đưa về kích thước 160 × 160.", "Recognizer biến crop thành vector 512 số.", "Vector được so sánh với các vector trong gallery của lớp.", "Hệ thống trả tên nếu đủ giống; nếu chưa chắc thì trả UNKNOWN."])
    h(doc, "1.10. Ví dụ về cosine similarity", 2)
    p(doc, "Giả sử vector query gần vector của sinh viên A hơn sinh viên B. Nếu điểm A là 0,78 và điểm B là 0,42, hệ thống có cơ sở để đề xuất A. Tuy nhiên nếu cả hai điểm chỉ là 0,51 và 0,50, kết quả quá gần nhau và nên chuyển sang ambiguous để giáo viên xem lại.")
    h(doc, "1.11. Tóm tắt chương", 2)
    p(doc, "Detector trả lời câu hỏi khuôn mặt nằm ở đâu. Recognizer trả lời khuôn mặt giống ai. ArcFace giúp học embedding có tính phân biệt. Threshold và review giúp biến dự đoán AI thành quy trình điểm danh an toàn hơn.")
    doc.add_page_break()

    h(doc, "Chương 2. Phân tích và thiết kế hệ thống")
    h(doc, "2.1. Người sử dụng", 2)
    table(doc, ["Đối tượng", "Nhu cầu"], [["Sinh viên", "Đăng ký khuôn mặt và xem lịch sử."], ["Giáo viên", "Chọn lớp, quét, xem và sửa kết quả."], ["Quản trị viên", "Quản lý tài khoản, lớp, camera và enrollment."], ["AI service", "Nhận ảnh, tạo vector và trả kết quả."]])
    h(doc, "2.2. Chức năng chính", 2)
    list_items(doc, ["Đăng nhập và phân quyền.", "Đăng ký khuôn mặt bằng nhiều frame.", "Quản lý lớp và roster.", "Quét điểm danh bằng camera.", "Nhận diện known student và UNKNOWN.", "Lưu attendance và evidence.", "Giáo viên chỉnh sửa kết quả."])
    h(doc, "2.3. Nguyên tắc thiết kế", 2)
    p(doc, "Frontend chỉ gọi Backend. Backend giữ logic nghiệp vụ và gọi AI service bằng request nội bộ có khóa bảo vệ. AI service xử lý ảnh và trả dữ liệu có cấu trúc. Database lưu lớp, enrollment, attendance và metadata của evidence.")
    figure(doc, architecture, "Hình 2.1. Kiến trúc thành phần chính")
    doc.add_page_break()

    h(doc, "Chương 2. Yêu cầu chức năng")
    h(doc, "2.3.1. Yêu cầu theo người dùng", 2)
    table(doc, ["Mã", "Yêu cầu", "Đối tượng"], [["FR-01", "Đăng nhập và phân quyền.", "Tất cả"], ["FR-02", "Tạo enrollment từ nhiều frame.", "Sinh viên / quản trị"], ["FR-03", "Tạo session điểm danh theo lớp.", "Giáo viên"], ["FR-04", "Nhận diện known và UNKNOWN.", "AI service"], ["FR-05", "Lưu kết quả và evidence.", "Backend"], ["FR-06", "Xem và sửa kết quả.", "Giáo viên"], ["FR-07", "Xem trạng thái model và service.", "Quản trị"]])
    h(doc, "2.3.2. Yêu cầu phi chức năng", 2)
    table(doc, ["Nhóm", "Yêu cầu cơ bản"], [["Dễ dùng", "Màn hình demo ít bước, thông báo rõ."], ["Tin cậy", "Không ghi attendance khi identity chưa chắc."], ["Bảo mật", "AI service không mở trực tiếp ra Internet."], ["Mở rộng", "Có thể thay detector hoặc recognizer riêng."], ["Theo dõi", "Có log request, score và lỗi tổng quát."]])
    h(doc, "2.3.3. Quy tắc nghiệp vụ", 2)
    p(doc, "Một sinh viên chỉ có một trạng thái cuối trong một session. Một identity ngoài roster không được ghi vào danh sách lớp. Trường hợp một khuôn mặt gắn với nhiều tài khoản cần bị từ chối ở bước enrollment.")
    doc.add_page_break()

    h(doc, "Chương 2. Luồng nghiệp vụ")
    h(doc, "2.4. Luồng điểm danh", 2)
    p(doc, "Giáo viên chọn lớp. Backend lấy sinh viên đã enrollment và nạp gallery cho session. Frontend chụp frame rồi gửi tới Backend. Backend chuyển frame cho AI service. Kết quả được lọc theo roster, loại UNKNOWN và tránh ghi trùng.")
    figure(doc, attendance, "Hình 2.2. Luồng điểm danh")
    h(doc, "2.5. Xử lý kết quả không chắc chắn", 2)
    p(doc, "Nếu điểm similarity thấp hơn threshold, hệ thống không gán identity. Nếu hai kết quả đầu quá gần nhau, hệ thống xem là ambiguous. Giáo viên có thể xem evidence và sửa trạng thái.")
    h(doc, "2.6. Enrollment", 2)
    p(doc, "Enrollment yêu cầu nhiều pose như nhìn thẳng, quay trái và quay phải. Các frame đạt chuẩn được tạo vector. Một khuôn mặt không được đăng ký cho nhiều tài khoản để giảm gian lận và lỗi dữ liệu.")
    doc.add_page_break()

    h(doc, "Chương 2. Use case và activity")
    h(doc, "2.6.1. Use case chính", 2)
    table(doc, ["Actor", "Use case", "Kết quả"], [["Giáo viên", "Mở session điểm danh", "Session có roster"], ["Giáo viên", "Quét camera", "Có danh sách prediction"], ["Giáo viên", "Review evidence", "Xác nhận hoặc sửa"], ["Sinh viên", "Enrollment", "Có gallery embedding"], ["Quản trị", "Quản lý lớp và user", "Dữ liệu sẵn sàng cho session"]])
    h(doc, "2.6.2. Activity đơn giản", 2)
    list_items(doc, ["Bắt đầu: giáo viên chọn lớp.", "Backend lấy roster và tạo session.", "Frontend gửi frame theo chu kỳ.", "AI service trả kết quả từng frame.", "Backend gộp kết quả và loại trùng.", "Giáo viên review các trường hợp cần chú ý.", "Kết thúc: session được lưu và có thể xuất báo cáo."])
    h(doc, "2.6.3. Trường hợp ngoại lệ", 2)
    p(doc, "Nếu camera không đọc được frame, giao diện báo lỗi kết nối. Nếu AI service không sẵn sàng, Backend không tạo attendance giả. Nếu score thấp, kết quả chuyển UNKNOWN và vẫn giữ evidence để review.")
    doc.add_page_break()

    h(doc, "Chương 2. Mô hình dữ liệu và API")
    h(doc, "2.7. Entity chính", 2)
    table(doc, ["Entity", "Vai trò"], [["User", "Tài khoản và role."], ["Section", "Môn học, lớp học phần và phòng."], ["Enrollment", "Trạng thái đăng ký khuôn mặt."], ["FaceEmbedding", "Vector đại diện khuôn mặt."], ["AttendanceSession", "Một lần giáo viên mở quét."], ["AttendanceRecord", "Trạng thái của sinh viên."], ["Evidence", "Ảnh hoặc đường dẫn ảnh minh chứng."]])
    h(doc, "2.8. API AI quan trọng", 2)
    table(doc, ["API", "Chức năng"], [["GET /health", "Kiểm tra service và model."], ["POST enrollment", "Nhận frame và tạo embedding."], ["PUT roster", "Nạp gallery theo session."], ["POST recognition", "Trả boxes, identity và score."], ["POST capture", "Nhận frame từ camera."]])
    h(doc, "2.9. Vì sao tách detector và recognizer", 2)
    p(doc, "Hai model có nhiệm vụ khác nhau nên dễ đánh giá và thay thế riêng. Khi thêm sinh viên, hệ thống chỉ cần tạo vector mới và cập nhật gallery; không cần train lại classifier theo tên.")
    doc.add_page_break()

    h(doc, "Chương 2. Thiết kế dữ liệu và luồng trạng thái")
    h(doc, "2.9.1. Dữ liệu cần lưu", 2)
    table(doc, ["Dữ liệu", "Ví dụ trường", "Lý do"], [["Roster", "student_id, name, status", "Xác định ai thuộc lớp"], ["Embedding", "vector 512D, created_at", "So khớp khuôn mặt"], ["Prediction", "identity, score, box", "Giải thích kết quả AI"], ["Session", "class_id, opened_at", "Nhóm các lần quét"], ["Evidence", "path, timestamp", "Kiểm tra lại sau này"]])
    h(doc, "2.9.2. Trạng thái attendance", 2)
    table(doc, ["Trạng thái", "Ý nghĩa"], [["PRESENT_AI", "AI đề xuất có mặt."], ["UNKNOWN", "Không đủ cơ sở nhận diện."], ["REVIEW", "Chờ giáo viên kiểm tra."], ["PRESENT_CONFIRMED", "Giáo viên xác nhận."], ["ABSENT", "Không thấy hoặc giáo viên đánh dấu vắng."]])
    h(doc, "2.9.3. Tính nhất quán", 2)
    p(doc, "Backend là nơi quyết định trạng thái cuối và chống ghi trùng. AI service chỉ trả prediction. Cách tách này giúp thay model mà không làm thay đổi quy tắc attendance trong hệ thống.")
    doc.add_page_break()

    h(doc, "Chương 3. Huấn luyện mô hình và tích hợp hệ thống")
    h(doc, "3.1. Quy trình chung", 2)
    p(doc, "Hai notebook có cùng flow: EDA, preprocessing, train, validation, xuất metric và inference ảnh thật. Khi đóng gói, cần giữ model, preprocessing, threshold và gallery format cùng nhau.")
    h(doc, "3.2. Dữ liệu cho YOLO11n", 2)
    p(doc, "Notebook detector dùng WIDER Face dạng YOLO. Có 12.880 ảnh train và 3.226 ảnh validation. Label chứa class và tọa độ bounding box được chuẩn hóa.")
    table(doc, ["Thông số", "YOLO11n"], [["Input", "640 × 640"], ["Epochs", "25"], ["Batch size", "8"], ["Train images", "12.880"], ["Validation images", "3.226"], ["Output", "Bounding box và confidence"]])
    h(doc, "3.3. Preprocessing detector", 2)
    p(doc, "Notebook kiểm tra file ảnh, đọc label, hiển thị ảnh trước xử lý và dùng letterbox để đưa ảnh về 640. Ảnh hỏng hoặc label sai cần được loại và ghi nhận.")
    doc.add_page_break()

    h(doc, "Chương 3. Chuẩn bị dữ liệu và augmentation")
    h(doc, "3.3.1. Kiểm tra dữ liệu", 2)
    table(doc, ["Bước", "Mục đích"], [["Đọc file ảnh", "Phát hiện file hỏng hoặc không mở được."], ["Kiểm tra label", "Đảm bảo box nằm trong ảnh."], ["Xem ảnh mẫu", "Phát hiện label lệch bằng mắt."], ["Tách train/validation", "Đo trên ảnh chưa dùng để cập nhật weight."], ["Ghi thông tin", "Có thể giải thích dataset khi report."]])
    h(doc, "3.3.2. Augmentation có ý nghĩa gì", 2)
    p(doc, "Các thay đổi nhẹ về scale, crop, màu và vị trí giúp detector gặp nhiều điều kiện hơn lúc train. Augmentation phải giữ đúng box. Không nên tạo biến đổi quá mạnh làm ảnh không giống bối cảnh camera thật.")
    h(doc, "3.3.3. Nguy cơ data leakage", 2)
    p(doc, "Nếu ảnh gần như giống nhau xuất hiện ở cả train và validation, metric có thể cao hơn thực tế. Vì vậy dữ liệu validation cần được giữ riêng và không dùng để chọn thủ công từng ảnh trong lúc train.")
    doc.add_page_break()

    h(doc, "Chương 3. Model nhận diện khuôn mặt")
    h(doc, "3.4. Dữ liệu cho FaceNet/ArcFace", 2)
    p(doc, "Notebook recognition bản latest dùng WebFace 112 × 112. Sau khi lọc, notebook chọn 1.500 identity, gồm 37.501 ảnh train và 9.383 ảnh validation. Quá trình train chạy 20 epoch và lưu checkpoint tốt nhất là `embedder_best.pt`.")
    h(doc, "3.5. Kiến trúc recognition", 2)
    table(doc, ["Thông số", "Giá trị"], [["Backbone", "InceptionResnetV1"], ["Input", "160 × 160"], ["Embedding", "512 chiều"], ["Training head", "ArcMarginProduct"], ["Tham số", f"{arcface_params:,}"], ["Epochs", str(arcface_epochs)], ["Optimizer", "AdamW"], ["Scheduler", "Cosine Annealing"]])
    h(doc, "3.6. Vì sao dùng 512 chiều", 2)
    p(doc, "512 chiều đủ để giữ thông tin nhận diện và vẫn dễ lưu, truyền và so sánh. Gallery của lớp chỉ chứa vector của sinh viên trong lớp, nên thời gian so sánh không phụ thuộc toàn bộ dataset train.")
    h(doc, "3.7. Loss và tối ưu", 2)
    p(doc, "Embedding và weight được chuẩn hóa trước khi tạo logits. Cross-entropy được tính trên logits ArcFace. Label smoothing giúp giảm việc model quá tự tin trên dữ liệu train. AdamW cập nhật weight và Cosine Annealing thay đổi learning rate theo epoch.")
    figure(doc, pipeline, "Hình 3.1. Chi tiết pipeline nhận diện")
    doc.add_page_break()

    h(doc, "Chương 3. Tham số huấn luyện ArcFace latest")
    h(doc, "3.7.1. Cấu hình đã chạy", 2)
    table(doc, ["Thông số", "Giá trị"], [["Dataset", "WebFace 112 × 112"], ["Identity được chọn", "1.500"], ["Train images", "37.501"], ["Validation images", "9.383"], ["Epochs", "20"], ["Embedding", "512 chiều"], ["Số tham số", "27.910.327"], ["GPU", "2 GPU trong notebook"], ["Checkpoint tốt nhất", "embedder_best.pt"]])
    h(doc, "3.7.2. Vì sao chọn các tham số này", 2)
    p(doc, "Ảnh input 160 × 160 phù hợp với preprocessing của FaceNet và giữ chi phí tính toán vừa phải. Embedding 512 chiều là đầu ra mà hệ thống gallery đang sử dụng. 20 epoch là cấu hình latest đã chạy; checkpoint tốt nhất được chọn theo validation Top-1, không chọn đơn giản theo epoch cuối.")
    h(doc, "3.7.3. Ý nghĩa ArcMarginProduct", 2)
    p(doc, "ArcMarginProduct tạo logits có thêm khoảng cách góc giữa các identity. Trong lúc train, head này giúp model học ranh giới rõ hơn. Khi deploy, hệ thống không cần lưu toàn bộ 1.500 class; chỉ cần embedding và gallery theo roster lớp.")
    doc.add_page_break()

    h(doc, "Chương 3. Tích hợp model vào hệ thống")
    h(doc, "3.8. AI service", 2)
    p(doc, "AI service chạy bằng FastAPI trong container riêng. Service load `face_best.pt` cho detector và `facenet_best.pt` cho recognizer. Hai weight nằm trong thư mục models và được mount read-only. Weight không đưa lên Git vì kích thước lớn.")
    h(doc, "3.9. Xử lý một frame", 2)
    list_items(doc, ["Decode và kiểm tra ảnh.", "YOLO tìm các khuôn mặt.", "Crop và alignment khuôn mặt.", "Recognizer tạo vector 512 chiều.", "So sánh vector với roster gallery.", "Áp dụng threshold và ambiguity margin.", "Trả identity, score, box và evidence crop."])
    h(doc, "3.10. Xử lý một session", 2)
    p(doc, "Frontend có thể gửi nhiều frame trong thời gian quét. Backend gom kết quả, loại trùng và lưu thời điểm xuất hiện đầu tiên. Khi đủ roster, hệ thống có thể dừng sớm để giảm tải.")
    h(doc, "3.11. An toàn khi tích hợp", 2)
    p(doc, "Frontend không gọi model trực tiếp. Backend kiểm tra JWT, role và khóa gọi AI service. AI service không nên mở trực tiếp cho người dùng ngoài hệ thống.")
    doc.add_page_break()

    h(doc, "Chương 3. Reproducibility và triển khai")
    h(doc, "3.11.1. Những thứ phải cố định", 2)
    table(doc, ["Thành phần", "Cần giữ giống nhau"], [["Model", "Checkpoint và kiến trúc."], ["Preprocessing", "Resize, normalize, alignment."], ["Gallery", "Format vector và identity mapping."], ["Threshold", "acceptThreshold và ambiguityMargin."], ["Version", "Notebook, package và commit code."]])
    h(doc, "3.11.2. Checklist chạy lại", 2)
    list_items(doc, ["Kiểm tra GPU hoặc CPU và phiên bản PyTorch.", "Đọc đúng dataset và seed nếu có.", "Lưu log từng epoch cùng checkpoint.", "Chạy validation sau khi train.", "Chạy inference trên ảnh ví dụ.", "Kiểm tra model load được trong AI service.", "Đối chiếu output notebook với output API."])
    h(doc, "3.11.3. Tại sao report ghi rõ checkpoint", 2)
    p(doc, "Hai file có thể cùng mang tên ArcFace nhưng được train ở thời điểm khác nhau. Ghi tên checkpoint và nguồn metric giúp người đọc biết con số nào thuộc notebook latest, con số nào thuộc model đang chạy trong hệ thống.")
    doc.add_page_break()

    h(doc, "Chương 3. Kết quả huấn luyện và metric")
    h(doc, "3.12. Metric YOLO11n", 2)
    p(doc, "Bảng sau lấy từ file metrics.json được sinh sau validation YOLO11n trên 3.226 ảnh và 39.707 instance khuôn mặt.")
    table(doc, ["Metric", "Giá trị", "Ý nghĩa"], [["Precision", f"{precision:.4f}", "Tỷ lệ dự đoán mặt đúng."], ["Recall", f"{recall:.4f}", "Tỷ lệ mặt thật được tìm thấy."], ["mAP50", f"{map50:.4f}", "Chất lượng box ở IoU 0.50."], ["mAP50-95", f"{map5095:.4f}", "Chất lượng box ở nhiều ngưỡng."]])
    h(doc, "3.13. Metric ArcFace từ notebook latest", 2)
    p(doc, "Đây là kết quả của notebook `face-attendance-arcface-recognition (1).ipynb` sau lần train mới. Validation được tính trên 9.383 ảnh của 1.500 identity. Phần gallery đánh giá trên cùng tập validation với gallery được tạo từ checkpoint tốt nhất; đây là metric của mô hình nhận diện, chưa phải metric production trên camera lớp học.")
    table(doc, ["Metric", "Giá trị", "Phạm vi"], [["Best validation Top-1", f"{arcface_val_top1:.4f}", "9.383 ảnh validation"], ["Validation Top-5", f"{arcface_val_top5:.4f}", "9.383 ảnh validation"], ["Gallery Top-1", f"{arcface_gallery_top1:.4f}", "1.500 identity"], ["Gallery Top-5", f"{arcface_gallery_top5:.4f}", "1.500 identity"], ["Mean true cosine", f"{arcface_true_cosine_mean:.4f}", "Query với identity đúng"], ["Mean margin", f"{arcface_margin_mean:.4f}", "Top-1 trừ Top-2"], ["Checkpoint", arcface_checkpoint, "Notebook latest"]])
    h(doc, "3.14. Metric model đang cài trong hệ thống", 2)
    p(doc, "Model `facenet_best.pt` trong thư mục models được kiểm tra thêm trên 40 frame enrollment của hệ thống, gồm 8 frame cho mỗi 5 sinh viên. Cách đo leave-one-out dùng một frame làm query và các frame còn lại làm gallery. Kết quả này chỉ là test nội bộ; không đại diện cho dữ liệu sinh viên mới và không nên gộp với metric notebook latest.")
    table(doc, ["Metric", "Giá trị", "Phạm vi"], [["Số identity", "5", "Dữ liệu enrollment hệ thống"], ["Số frame", "40", "8 frame / identity"], ["Cách đo", "Leave-one-out", "Mỗi query bỏ khỏi gallery"], ["Top-1", f"{system_top1:.4f}", "40 query nội bộ"], ["Top-5", f"{system_top5:.4f}", "40 query nội bộ"], ["Mean cosine", f"{system_mean_cosine:.4f}", "Query với gallery cùng identity"], ["Min cosine", f"{system_min_cosine:.4f}", "Giá trị thấp nhất"], ["Mean margin", f"{system_mean_margin:.4f}", "Top-1 trừ Top-2"]])
    h(doc, "3.15. Cách chốt metric cuối", 2)
    list_items(doc, [f"Notebook latest đạt validation Top-1 {arcface_val_top1:.4f} và Top-5 {arcface_val_top5:.4f}.", f"Gallery validation đạt Top-1 {arcface_gallery_top1:.4f} và Top-5 {arcface_gallery_top5:.4f}.", f"Model đang cài trong hệ thống đạt leave-one-out Top-1 {system_top1:.4f} trên 40 frame enrollment.", "Hai phép đo dùng dữ liệu và checkpoint khác nhau nên phải ghi tách riêng.", "Bổ sung test người ngoài roster và test camera thật trước khi kết luận production.", "Không dùng accuracy của training làm accuracy production."])
    doc.add_page_break()

    h(doc, "Chương 3. Cách đọc kết quả metric")
    h(doc, "3.16. Validation và gallery không giống nhau", 2)
    p(doc, "Validation Top-1 đo identity đúng đứng đầu trên các query validation. Gallery Top-1 mô phỏng cách query được so với gallery của các identity. Hai số có thể khác nhau vì cách tạo gallery và cách tổng hợp kết quả khác nhau.")
    table(doc, ["Kết quả latest", "Giá trị", "Diễn giải"], [["Validation Top-1", f"{arcface_val_top1:.4f}", "Khoảng 89/100 query đúng ở vị trí đầu."], ["Validation Top-5", f"{arcface_val_top5:.4f}", "Khoảng 92/100 query có đáp án đúng trong 5 vị trí đầu."], ["Gallery Top-1", f"{arcface_gallery_top1:.4f}", "Kết quả khi dùng gallery trên validation."], ["Mean true cosine", f"{arcface_true_cosine_mean:.4f}", "Mức giống trung bình với identity đúng."], ["Mean margin", f"{arcface_margin_mean:.4f}", "Khoảng cách trung bình giữa đáp án đầu và thứ hai."]])
    h(doc, "3.17. Không so sánh trực tiếp với metric detector", 2)
    p(doc, "Precision, Recall và mAP của YOLO đo vị trí khuôn mặt. Top-1 và Top-5 của ArcFace đo identity. Một pipeline có thể tìm đúng box nhưng nhận sai tên, hoặc nhận tốt khi box đã được crop đúng. Vì vậy report giữ hai nhóm metric riêng.")
    h(doc, "3.18. Việc cần làm để chốt", 2)
    p(doc, "Cần có một tập test camera thật với identity known, người ngoài roster, nhiều góc mặt và nhiều điều kiện sáng. Tập test này phải giữ độc lập với tập enrollment và validation. Sau đó mới chọn threshold, tính false accept/false reject và kết luận khả năng dùng thực tế.")
    doc.add_page_break()

    h(doc, "Chương 4. Kết quả, kiểm thử và đánh giá hệ thống")
    h(doc, "4.1. Môi trường", 2)
    table(doc, ["Thành phần", "Cấu hình"], [["Training", "Kaggle notebook có GPU."], ["Serving", "Docker Compose."], ["Detector", "YOLO11n với face_best.pt."], ["Recognizer", "InceptionResnetV1 với facenet_best.pt."], ["Input", "Webcam, ảnh upload hoặc camera frame."], ["Output", "Identity, score, attendance và evidence."]])
    h(doc, "4.2. Kết quả detector", 2)
    p(doc, f"Validation detector đạt Precision {precision:.4f}, Recall {recall:.4f}, mAP50 {map50:.4f} và mAP50-95 {map5095:.4f}. Kết quả cho thấy model đã học được khả năng tìm khuôn mặt, nhưng Recall thấp hơn Precision nên vẫn có trường hợp bỏ sót.")
    p(doc, "Log validation có cảnh báo max_det do một số ảnh chứa rất nhiều object. Vì vậy cần chạy lại với giới hạn phù hợp trước khi dùng bảng này làm kết quả benchmark cuối.")
    h(doc, "4.3. Kết quả recognition", 2)
    p(doc, f"Recognition cần báo cáo riêng vì detector đúng không có nghĩa identity đúng. Notebook ArcFace latest đạt validation Top-1 {arcface_val_top1:.4f}, Top-5 {arcface_val_top5:.4f}; gallery validation đạt Top-1 {arcface_gallery_top1:.4f}. Model đang tích hợp trong hệ thống có phép đo nội bộ leave-one-out Top-1 {system_top1:.4f} trên 40 frame của 5 identity. Cần test thêm người ngoài roster trước khi kết luận production.")
    h(doc, "4.4. Kiểm thử chức năng", 2)
    table(doc, ["ID", "Tình huống", "Kỳ vọng"], [["T01", "Đăng nhập", "Đúng màn hình theo role."], ["T02", "Enrollment", "Lưu vector sau khi đủ frame."], ["T03", "Mở camera", "Session bắt đầu quét."], ["T04", "Known student", "Ghi đúng identity."], ["T05", "Unknown", "Không ghi nhầm roster."], ["T06", "Trùng khuôn mặt", "Từ chối và báo lỗi."], ["T07", "Evidence", "Giáo viên xem được ảnh."], ["T08", "Chỉnh sửa", "Trạng thái mới được lưu."]])
    doc.add_page_break()

    h(doc, "Chương 4. Phương pháp kiểm thử")
    h(doc, "4.4.1. Kiểm thử model", 2)
    table(doc, ["Nhóm test", "Dữ liệu", "Metric"], [["Detector validation", "3.226 ảnh, 39.707 instance", "Precision, Recall, mAP"], ["ArcFace validation", "9.383 ảnh, 1.500 identity", "Top-1, Top-5"], ["System diagnostic", "40 enrollment frame, 5 identity", "Leave-one-out Top-1/Top-5"], ["Unknown test", "Người ngoài roster", "False accept rate"], ["Camera test", "Ảnh/video lớp học thật", "Accuracy, latency, failure case"]])
    h(doc, "4.4.2. Kiểm thử theo tầng", 2)
    list_items(doc, ["Tầng model: kiểm tra output box, vector và score.", "Tầng AI service: kiểm tra API, model load và lỗi ảnh.", "Tầng Backend: kiểm tra roster, chống trùng và lưu session.", "Tầng Frontend: kiểm tra camera, hiển thị và teacher review.", "Tầng end-to-end: chạy một session từ đăng nhập đến kết thúc."])
    h(doc, "4.4.3. Tiêu chí Pass/Fail", 2)
    p(doc, "Test Pass khi kết quả đúng với kỳ vọng và không tạo side effect sai. Test Fail khi service trả lỗi không rõ, ghi nhầm roster, mất evidence hoặc không thể khôi phục session. Các case Fail cần ghi lại input, log và bước tái hiện.")
    doc.add_page_break()

    h(doc, "Chương 4. Phân tích lỗi và giới hạn")
    h(doc, "4.5. Các lỗi có thể gặp", 2)
    table(doc, ["Tình huống", "Nguyên nhân", "Cách giảm lỗi"], [["Ảnh mờ", "Camera rung hoặc người di chuyển.", "Quét nhiều frame."], ["Ánh sáng yếu", "Mặt có ít thông tin.", "Bổ sung dữ liệu low-light."], ["Góc nghiêng", "Khác ảnh enrollment.", "Enrollment nhiều pose."], ["Mặt nhỏ", "Khoảng cách xa.", "Điều chỉnh camera."], ["Người giống nhau", "Vector gần nhau.", "Threshold và teacher review."], ["Ảnh replay", "Ảnh/video giả.", "Bổ sung liveness."]])
    h(doc, "4.6. Giới hạn", 2)
    p(doc, "Metric public không đại diện hoàn toàn cho camera lớp học. Threshold có thể thay đổi theo ánh sáng, vị trí camera và chất lượng enrollment. Model hiện tại chưa đủ để kết luận chống giả mạo nếu chưa có liveness test riêng.")
    h(doc, "4.7. Đánh giá MVP", 2)
    p(doc, "SPAS phù hợp làm prototype AI end-to-end và demo. Hệ thống đã thể hiện được kết nối từ camera, detector, recognizer, gallery, roster đến attendance. Để triển khai rộng cần thêm dữ liệu camera thật, calibration, liveness và test tải.")
    doc.add_page_break()

    h(doc, "Chương 4. Đánh giá vận hành và quyền riêng tư")
    h(doc, "4.7.1. Các yếu tố ảnh hưởng tốc độ", 2)
    table(doc, ["Yếu tố", "Ảnh hưởng", "Cách xử lý"], [["Số khuôn mặt", "Nhiều crop cần nhiều lần embedding.", "Batching hoặc tracking."], ["Độ phân giải", "Ảnh lớn tốn thời gian decode.", "Giới hạn kích thước hợp lý."], ["Khoảng cách gallery", "Roster lớn tăng số so sánh.", "Chỉ nạp roster của lớp."], ["Mạng nội bộ", "Frame chậm tới AI service.", "Giảm kích thước hoặc queue."], ["Lưu evidence", "I/O tăng khi lưu nhiều ảnh.", "Chính sách lưu và nén ảnh."]])
    h(doc, "4.7.2. Quyền riêng tư trong demo", 2)
    list_items(doc, ["Chỉ thu thập khi người dùng biết mục đích.", "Chỉ cho role phù hợp xem ảnh và embedding.", "Không hiển thị vector thô trên giao diện.", "Không đưa ảnh khuôn mặt vào log lỗi.", "Có thời hạn lưu và cách xóa dữ liệu.", "Cho phép giáo viên sửa nhưng vẫn giữ audit cơ bản."])
    h(doc, "4.7.3. Trách nhiệm của người vận hành", 2)
    p(doc, "Giáo viên cần kiểm tra các case UNKNOWN hoặc ambiguous và không dùng một metric duy nhất để quyết định kỷ luật. Hệ thống nên được xem là công cụ hỗ trợ, đặc biệt khi dữ liệu camera có thể bị mờ hoặc thiếu sáng.")
    doc.add_page_break()

    h(doc, "Chương 4. Kịch bản demo và tiêu chí nghiệm thu")
    h(doc, "4.8. Kịch bản demo", 2)
    list_items(doc, ["Đăng nhập tài khoản giáo viên.", "Chọn lớp và xem roster.", "Mở camera và quét frame.", "Quan sát box, identity và score.", "Kiểm tra UNKNOWN bằng người ngoài roster.", "Mở evidence.", "Chỉnh trạng thái một sinh viên."])
    h(doc, "4.9. So sánh với điểm danh thủ công", 2)
    table(doc, ["Tiêu chí", "Thủ công", "SPAS"], [["Ghi nhận", "Đọc tên và nhập tay.", "Camera và AI tạo kết quả."], ["Bằng chứng", "Thường không có ảnh.", "Có evidence."], ["Ngoại lệ", "Giáo viên tự sửa.", "Có teacher review."], ["Rủi ro", "Nhập nhầm hoặc bỏ sót.", "Phụ thuộc camera và threshold."]])
    h(doc, "4.10. Tiêu chí nghiệm thu", 2)
    list_items(doc, ["Role đăng nhập đúng.", "Enrollment tạo được vector.", "Known student được ghi nhận.", "Unknown không bị gán nhầm.", "Evidence xem lại được.", "Giáo viên sửa được kết quả.", "Metric và giới hạn được báo cáo trung thực."])
    doc.add_page_break()

    h(doc, "Chương 4. Bài học từ quá trình làm dự án")
    h(doc, "4.11. Bài học về dữ liệu", 2)
    p(doc, "Metric chỉ đáng tin khi biết rõ metric được đo trên dữ liệu nào. Dữ liệu training, validation, enrollment và camera thật có mục đích khác nhau. Report phải ghi số ảnh, số identity, cách chia tập và phạm vi kết luận.")
    h(doc, "4.12. Bài học về AI", 2)
    p(doc, "Không nên đánh giá nhận diện chỉ bằng một ví dụ đúng. Cần đo theo query, identity, threshold và trường hợp unknown. Detector và recognizer nên được đánh giá riêng để biết lỗi nằm ở bước nào.")
    h(doc, "4.13. Bài học về sản phẩm", 2)
    p(doc, "Một model tốt vẫn cần API, gallery, roster, evidence, quyền truy cập và màn hình review. Teacher review không phải phần phụ; đó là cơ chế xử lý các trường hợp AI chưa chắc chắn.")
    h(doc, "4.14. Bài học về trình bày", 2)
    p(doc, "Khi thuyết trình, nên bắt đầu từ vấn đề và demo pipeline, sau đó mới giải thích layer, embedding và metric. Các con số cần đi cùng câu giải thích đơn giản để cả team và giáo viên đều theo dõi được.")
    doc.add_page_break()

    h(doc, "Kết luận và hướng phát triển")
    p(doc, "Nhóm đã xây dựng prototype điểm danh tự động có luồng từ ảnh camera đến attendance record. AI gồm detector YOLO11n, recognizer tạo embedding 512 chiều, cosine similarity và gallery theo roster. Hệ thống web gồm Frontend, Backend, AI service, database và evidence.")
    p(doc, "Thiết kế không để AI tự quyết định tuyệt đối. Khi score thấp hoặc kết quả không rõ, hệ thống trả UNKNOWN để giáo viên kiểm tra. Điều này giảm rủi ro ghi nhầm người.")
    h(doc, "Hướng phát triển", 2)
    list_items(doc, ["Hoàn tất metric recognition trên test độc lập.", "Chạy lại YOLO validation với max_det phù hợp.", "Calibration threshold theo camera.", "Bổ sung dữ liệu lớp học có consent.", "Thêm liveness và anti-spoofing.", "Tối ưu tracking, batching và nhiều camera.", "Bổ sung retention policy và quy trình xóa dữ liệu."])
    doc.add_page_break()

    h(doc, "Tài liệu tham khảo")
    references = [
        "Samsung Innovation Campus. AI Course Capstone Project Final Report Template.",
        "Samsung Innovation Campus. AI Course Capstone Project Action Plan Template.",
        "Ultralytics. YOLO Documentation. https://docs.ultralytics.com/",
        "Deng, J. et al. ArcFace: Additive Angular Margin Loss for Deep Face Recognition. CVPR, 2019.",
        "Schroff, F. et al. FaceNet: A Unified Embedding for Face Recognition and Clustering. CVPR, 2015.",
        "SPAS repository. https://github.com/NguyenVanTruong123/SIC-Smart-Attendance",
        "Notebook: face-attendance-yolo11n-detector.ipynb.",
        "Notebook: face-attendance-arcface-recognition (1).ipynb (latest training output).",
    ]
    for index, reference in enumerate(references, 1):
        doc.add_paragraph(f"[{index}] {reference}")
    doc.add_page_break()

    h(doc, "Phụ lục. Bảng số liệu cần hoàn tất")
    p(doc, "Các số liệu dưới đây tách thành hai nguồn: notebook latest và model đang cài trong hệ thống. Cần bổ sung test độc lập trước khi gọi đây là kết quả production.")
    table(doc, ["Mục", "Giá trị", "Nguồn"], [["Latest validation Top-1", f"{arcface_val_top1:.4f}", "Notebook, 9.383 ảnh"], ["Latest validation Top-5", f"{arcface_val_top5:.4f}", "Notebook, 9.383 ảnh"], ["Latest gallery Top-1", f"{arcface_gallery_top1:.4f}", "Notebook, 1.500 identity"], ["Latest gallery Top-5", f"{arcface_gallery_top5:.4f}", "Notebook, 1.500 identity"], ["System leave-one-out Top-1", f"{system_top1:.4f}", "40 enrollment frame"], ["System leave-one-out Top-5", f"{system_top5:.4f}", "40 enrollment frame"], ["System checkpoint", "facenet_best.pt", "models/"], ["Test độc lập người mới", "Cần bổ sung", "Dữ liệu camera có label"]])
    h(doc, "Checklist trước khi nộp", 2)
    list_items(doc, ["Cập nhật metric recognition.", "Chốt metric detector sau khi xử lý cảnh báo max_det.", "Chèn ảnh EDA, preprocessing và inference.", "Đánh dấu Pass/Fail cho test case.", "Kiểm tra mục lục, số hình, số bảng và tài liệu tham khảo."])
    # Normalize the visible heading hierarchy after all content is written.
    # This keeps the report body readable while making the Samsung-style TOC
    # reflect the requested five-chapter structure.
    chapter_headings = {
        "Chương 1. Cơ sở lý thuyết": (2, True, "Chương 2. Cơ sở lý thuyết", None),
        "Chương 1. Cơ sở lý thuyết tiếp theo": (2, False, "2.5. Các khái niệm bổ sung", "2.5"),
        "Chương 1. Ví dụ minh họa dễ hiểu": (2, False, "2.6. Ví dụ minh họa dễ hiểu", "2.6"),
        "Chương 2. Phân tích và thiết kế hệ thống": (3, True, "Chương 3. Phân tích và thiết kế hệ thống", None),
        "Chương 2. Yêu cầu chức năng": (3, False, "3.4. Yêu cầu chức năng", "3.4"),
        "Chương 2. Luồng nghiệp vụ": (3, False, "3.5. Luồng nghiệp vụ", "3.5"),
        "Chương 2. Use case và activity": (3, False, "3.6. Use case và activity", "3.6"),
        "Chương 2. Mô hình dữ liệu và API": (3, False, "3.7. Mô hình dữ liệu và API", "3.7"),
        "Chương 2. Thiết kế dữ liệu và luồng trạng thái": (3, False, "3.8. Thiết kế dữ liệu và luồng trạng thái", "3.8"),
        "Chương 3. Huấn luyện mô hình và tích hợp hệ thống": (4, True, "Chương 4. Huấn luyện mô hình và tích hợp hệ thống", None),
        "Chương 3. Chuẩn bị dữ liệu và augmentation": (4, False, "4.4. Chuẩn bị dữ liệu và augmentation", "4.4"),
        "Chương 3. Model nhận diện khuôn mặt": (4, False, "4.5. Model nhận diện khuôn mặt", "4.5"),
        "Chương 3. Tham số huấn luyện ArcFace latest": (4, False, "4.6. Tham số huấn luyện ArcFace latest", "4.6"),
        "Chương 3. Tích hợp model vào hệ thống": (4, False, "4.7. Tích hợp model vào hệ thống", "4.7"),
        "Chương 3. Reproducibility và triển khai": (4, False, "4.8. Reproducibility và triển khai", "4.8"),
        "Chương 3. Kết quả huấn luyện và metric": (4, False, "4.9. Kết quả huấn luyện và metric", "4.9"),
        "Chương 3. Cách đọc kết quả metric": (4, False, "4.10. Cách đọc kết quả metric", "4.10"),
        "Chương 4. Kết quả, kiểm thử và đánh giá hệ thống": (5, True, "Chương 5. Kết quả, kiểm thử và đánh giá hệ thống", None),
        "Chương 4. Phương pháp kiểm thử": (5, False, "5.5. Phương pháp kiểm thử", "5.5"),
        "Chương 4. Phân tích lỗi và giới hạn": (5, False, "5.6. Phân tích lỗi và giới hạn", "5.6"),
        "Chương 4. Đánh giá vận hành và quyền riêng tư": (5, False, "5.7. Đánh giá vận hành và quyền riêng tư", "5.7"),
        "Chương 4. Kịch bản demo và tiêu chí nghiệm thu": (5, False, "5.8. Kịch bản demo và tiêu chí nghiệm thu", "5.8"),
        "Chương 4. Bài học từ quá trình làm dự án": (5, False, "5.9. Bài học từ quá trình làm dự án", "5.9"),
    }
    problem_headings = {
        "1. Bối cảnh": "1.1. Bối cảnh",
        "2. Mục tiêu": "1.2. Mục tiêu",
        "3. Phạm vi": "1.3. Phạm vi",
        "4. Vì sao cần AI hỗ trợ": "1.5. Vì sao cần AI hỗ trợ",
        "5. Câu hỏi AI của dự án": "1.6. Câu hỏi AI của dự án",
        "6. Kết quả mong muốn": "1.7. Kết quả mong muốn",
    }
    figure_captions = {
        "Hình 1.1. Pipeline AI từ ảnh đến identity": "Hình 2.1. Pipeline AI từ ảnh đến identity",
        "Hình 2.1. Kiến trúc thành phần chính": "Hình 3.1. Kiến trúc thành phần chính",
        "Hình 2.2. Luồng điểm danh": "Hình 3.2. Luồng điểm danh",
        "Hình 3.1. Chi tiết pipeline nhận diện": "Hình 4.1. Chi tiết pipeline nhận diện",
    }
    current_chapter = 1
    subsection_prefix = None
    subsection_child_number = 0
    overview_mode = False
    for paragraph in doc.paragraphs:
        original = paragraph.text.strip()
        if original == "Đặt vấn đề":
            set_text = "Chương 1. Đặt vấn đề và tổng quan hệ thống"
            paragraph.style = doc.styles["Heading 1"]
            if paragraph.runs:
                paragraph.runs[0].text = set_text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(set_text)
            current_chapter = 1
            subsection_prefix = None
            subsection_child_number = 0
            overview_mode = False
            continue
        if original == "Đặt vấn đề. Cách tiếp cận của nhóm":
            paragraph.style = doc.styles["Heading 2"]
            text = "1.4. Tổng quan hệ thống"
            if paragraph.runs:
                paragraph.runs[0].text = text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(text)
            subsection_prefix = None
            subsection_child_number = 0
            overview_mode = True
            continue
        if original in chapter_headings:
            target_chapter, is_main, replacement, subsection_prefix = chapter_headings[original]
            current_chapter = target_chapter
            subsection_child_number = 0
            overview_mode = False
            paragraph.style = doc.styles["Heading 1" if is_main else "Heading 2"]
            if paragraph.runs:
                paragraph.runs[0].text = replacement
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(replacement)
            continue
        if original in problem_headings:
            replacement = problem_headings[original]
            if paragraph.runs:
                paragraph.runs[0].text = replacement
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(replacement)
            continue
        if original in figure_captions:
            replacement = figure_captions[original]
            if paragraph.runs:
                paragraph.runs[0].text = replacement
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(replacement)
            continue
        if paragraph.style and paragraph.style.name == "Heading 2":
            match = re.match(r"^(\d+(?:\.\d+)+)\.\s+(.+)$", original)
            if match:
                old_number, title = match.groups()
                if overview_mode:
                    old_first = old_number.split(".", 1)[0]
                    overview_number = {"4": "1.5", "5": "1.6", "6": "1.7"}.get(old_first)
                    replacement = f"{overview_number}. {title}" if overview_number else original
                elif subsection_prefix:
                    subsection_child_number += 1
                    replacement = f"{subsection_prefix}.{subsection_child_number}. {title}"
                else:
                    suffix = old_number.split(".", 1)[1]
                    replacement = f"{current_chapter}.{suffix}. {title}"
                if replacement != original:
                    if paragraph.runs:
                        paragraph.runs[0].text = replacement
                        for run in paragraph.runs[1:]:
                            run.text = ""
                    else:
                        paragraph.add_run(replacement)

    output = OUT / "SPAS_Final_Report_Vietnamese_updated.docx"
    doc.save(output)
    return output


def clone_slide(prs: Presentation, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in source.shapes:
        slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return slide


def replace_text(shape, mapping: dict[str, str]) -> None:
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            replace_text(child, mapping)
    if not hasattr(shape, "text_frame"):
        return
    original = shape.text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            for old, new in mapping.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)
    # Some template placeholders are split across multiple runs. Handle those
    # after the run-preserving pass so they are not left visible in the deck.
    updated = shape.text
    for old, new in mapping.items():
        updated = updated.replace(old, new)
    if updated != original and updated != shape.text:
        shape.text = updated


def build_slides() -> Path:
    prs = Presentation(BASE / "SIC_AI_Capstone Project_Presentation Slide Template.pptx")
    arcface_val_top1 = 0.8895875519556645
    arcface_val_top5 = 0.9208142385164659
    arcface_gallery_top1 = 0.8801023126931685
    system_top1 = 0.975
    system_top5 = 1.0
    system_mean_cosine = 0.7936312291771174
    replace_text(prs.slides[0], {"Team Name": "SIC AI Team", "Project\vName": "Smart Passive Attendance System"})
    replace_text(prs.slides[1], {"Project Name": "Smart Passive Attendance System"})
    template_slide = prs.slides[2]
    slide_content = [
        ("Bài toán", ["Điểm danh thủ công mất thời gian", "Khó kiểm tra khi có tranh chấp", "Mục tiêu: AI tạo kết quả và ảnh minh chứng"]),
        ("Mục tiêu MVP", ["Đăng ký khuôn mặt", "Điểm danh theo roster lớp", "Từ chối UNKNOWN", "Giáo viên xem và sửa kết quả"]),
        ("Kiến trúc hệ thống", ["Frontend gọi Backend", "Backend gọi AI service nội bộ", "AI chạy YOLO11n và FaceNet 512D", "Backend lưu attendance và evidence"]),
        ("Pipeline AI", ["Ảnh camera", "YOLO11n tìm bounding box", "Crop và chuẩn hóa", "Embedding 512 chiều", "Cosine similarity với gallery"]),
        ("Model 1: YOLO11n", ["Input: 640 × 640", "25 epoch, batch 8", "Output: box và confidence", "Precision, Recall, mAP50, mAP50-95"]),
        ("Model 2: FaceNet và ArcFace", ["Input: 160 × 160", "InceptionResnetV1", "Embedding: 512 chiều", "ArcMarginProduct khi train", "Cosine similarity khi deploy"]),
        ("Dữ liệu", ["YOLO: 12.880 train, 3.226 validation", "ArcFace: 1.500 identity", "Kiểm tra ảnh và label", "Có EDA trước training"]),
        ("Metric YOLO đã có", [f"Precision: {0.8253197803081475:.4f}", f"Recall: {0.5617649280983202:.4f}", f"mAP50: {0.6437364126984133:.4f}", f"mAP50-95: {0.33972672829922285:.4f}"]),
        ("Metric ArcFace latest", [f"Validation Top-1: {arcface_val_top1:.4f}", f"Validation Top-5: {arcface_val_top5:.4f}", f"Gallery Top-1: {arcface_gallery_top1:.4f}", "1.500 identity; 9.383 ảnh validation"]),
        ("Metric model hệ thống", ["Dùng facenet_best.pt", f"Leave-one-out Top-1: {system_top1:.4f}", f"Top-5: {system_top5:.4f}; mean cosine: {system_mean_cosine:.4f}", "Test nội bộ trên 40 enrollment frame"]),
        ("Luồng demo", ["GV chọn lớp", "Mở camera và quét frame", "AI trả identity và score", "Backend ghi attendance", "GV xem evidence"]),
        ("Kiểm thử chính", ["Đăng nhập và phân quyền", "Enrollment", "Known student", "UNKNOWN", "Trùng khuôn mặt", "Evidence và chỉnh sửa"]),
        ("Giới hạn", ["Phụ thuộc ánh sáng, pose, blur và camera", "Threshold cần calibration", "Chưa kết luận về liveness", "Chưa đánh giá nhiều camera"]),
        ("Hướng phát triển", ["Hoàn tất metric recognition độc lập", "Bổ sung dữ liệu camera thật", "Liveness và anti-spoofing", "Tracking, batching, multi-camera"]),
        ("Kết luận", ["SPAS là pipeline AI end-to-end", "Detector và recognizer đánh giá riêng", "Kết quả vẫn có teacher review", "Model có thể thay thế độc lập"]),
    ]
    for index, (title, lines) in enumerate(slide_content, 1):
        slide = clone_slide(prs, template_slide)
        replace_text(slide, {"Slide Title": title, "Subtitle": "SPAS - Smart Passive Attendance System", "1.1": f"{index:02d}", "01": f"{index:02d}", "UNIT": "AI COURSE", "Level-1\nLevel-2": "\n".join(lines), "Please edit project name and page number in slide master.": ""})
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[2])
    # The supplied template contains one intentionally blank trailing slide.
    if len(prs.slides) > 2 and not any(sh.text.strip() for sh in prs.slides[2].shapes if hasattr(sh, "text")):
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[2])
    output = OUT / "SPAS_Presentation_Vietnamese.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    report = build_report()
    slides = build_slides()
    print(json.dumps({"report": str(report), "slides": str(slides)}, ensure_ascii=False, indent=2))
