from docx import Document
import json
d=Document(r"deliverables/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO.docx")
out={"paragraphs":[{"i":i,"text":p.text.strip()} for i,p in enumerate(d.paragraphs) if p.text.strip()],"tables":[{"i":i,"head":[c.text for c in t.rows[0].cells] if t.rows else []} for i,t in enumerate(d.tables)]}
open("tmp/report_inspection.json","w",encoding="utf-8").write(json.dumps(out,ensure_ascii=False,indent=2))
