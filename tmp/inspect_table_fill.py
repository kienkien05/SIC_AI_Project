from docx import Document
from docx.oxml.ns import qn
from collections import Counter
import json
d=Document(r"tmp/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO_final.docx")
fills=Counter()
for t in d.tables:
 for c in t.rows[0].cells:
  tcPr=c._tc.tcPr
  sh=tcPr.find(qn("w:shd"))
  fills[sh.get(qn("w:fill")) if sh is not None else "none"]+=1
out={"table_count":len(d.tables),"header_fill_counts":dict(fills),"tables_styles":sorted(set(t.style.name if t.style else "" for t in d.tables))}
open(r"tmp/table_fill_info.json","w",encoding="utf-8").write(json.dumps(out,ensure_ascii=False,indent=2))
