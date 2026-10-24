from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json,re
f=r"tmp/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO_refreshed.docx"
d=Document(f)
errors=[];runs=0
for ti,t in enumerate(d.tables):
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   for p in c.paragraphs:
    if p.text.strip():
     for r in p.runs:
      if r.text:
       runs+=1
       if r.font.name!="Times New Roman" or not r.font.size or abs(r.font.size.pt-13)>0.01: errors.append({"table":ti,"row":ri,"cell":ci,"text":r.text[:25],"font":r.font.name,"size":r.font.size.pt if r.font.size else None})
expected_meta="Phần II trình bày kiến thức và quy trình gắn với hai nhiệm vụ này, tập trung vào cách thực hiện và bài học chuyên môn."
meta_present=any(p.text.strip()==expected_meta for p in d.paragraphs)
page_break_chapters=[]
for i,p in enumerate(d.paragraphs):
 if p.text.strip().startswith("CHƯƠNG ") and p.paragraph_format.page_break_before:
  page_break_chapters.append(p.text.strip())
 if not p.text.strip() and any(b.get(qn("w:type"))=="page" for b in p._p.iter(qn("w:br"))):
  nxt=next((q.text.strip() for q in d.paragraphs[i+1:] if q.text.strip()),"")
  if nxt.startswith("CHƯƠNG "):page_break_chapters.append(nxt)
fig_captions=[p for p in d.paragraphs if p.text.strip().startswith(("Hình 1.1. Sơ đồ tổ chức","Hình 1.2. Cơ cấu tổ chức")) and "[11]" in p.text]
tbl_caps=[p for p in d.paragraphs if p.text.strip().startswith("Bảng ") and p.text.strip().split()[1][0].isdigit()]
out={"table_count":len(d.tables),"formatted_runs":runs,"font_violations":errors[:30],"font_violation_count":len(errors),"chapter_forced_breaks":page_break_chapters,"meta_line_present":meta_present,"table_caption_paragraphs":len(tbl_caps),"figure_caption_paragraphs":len(fig_captions),"figure_captions_italic_centered":[{"text":p.text.strip(),"italic":all(r.italic for r in p.runs if r.text.strip()),"centered":p.alignment==WD_ALIGN_PARAGRAPH.CENTER} for p in fig_captions]}
open(r"tmp/format_final_validation.json","w",encoding="utf-8").write(json.dumps(out,ensure_ascii=False,indent=2))
