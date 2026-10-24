from docx import Document
from pathlib import Path
import sys
src,out=map(Path,sys.argv[1:3])
d=Document(src)
old='Mẫu báo cáo yêu cầu giới thiệu hoạt động của đơn vị trong hai năm gần nhất. Tài liệu hội nhập hiện có giới thiệu các nhóm sản phẩm và dịch vụ nhưng không phải báo cáo tài chính và không cung cấp số liệu kinh doanh theo từng năm. Vì vậy, báo cáo ghi rõ giới hạn của nguồn thay vì tự suy đoán hoặc dùng số liệu tổng quát thay cho kết quả từng năm.'
new='Mẫu báo cáo yêu cầu giới thiệu hoạt động của đơn vị trong hai năm gần nhất. Tài liệu hội nhập không phải báo cáo tài chính nên không có số liệu doanh thu theo từng năm. Phần này bổ sung thông tin pháp lý và dữ liệu công khai; số liệu ước tính hoặc chưa được công bố được ghi rõ.'
p=next(p for p in d.paragraphs if p.text==old)
p.runs[0].text=new
for r in p.runs[1:]:r.text=''
h=next(p for p in d.paragraphs if p.text=='1.5. Tình hình hoạt động trong hai năm gần nhất')
h.paragraph_format.page_break_before=True
d.save(out)
