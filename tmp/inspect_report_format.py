from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
d=Document(r"deliverables/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO.docx")
paras=[]
for i,p in enumerate(d.paragraphs):
 t=p.text.strip()
 if t and (t.startswith("PHẦN ") or t.startswith("CHƯƠNG ") or t.startswith("KẾT LUẬN") or "Tài liệu tham khảo" in t or "KẾT LUẬN" in t):
  ppr=p._p.find(qn("w:pPr"))
  page_before=bool(ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None)
  breaks=sum(1 for b in p._p.iter(qn("w:br")) if b.get(qn("w:type"))=="page")
  paras.append({"i":i,"text":t,"style":p.style.name,"pageBefore":page_before,"explicitPageBreaks":breaks,"keepNext":bool(ppr is not None and ppr.find(qn("w:keepNext")) is not None)})
tabs=[]
for ti,t in enumerate(d.tables):
 rows=[]
 for ri,row in enumerate(t.rows):
  cells=[]
  for ci,c in enumerate(row.cells):
   pf=[]
   for p in c.paragraphs:
    runs=[{"font":r.font.name,"size":r.font.size.pt if r.font.size else None,"bold":r.bold,"italic":r.italic} for r in p.runs if r.text.strip()]
    if p.text.strip():pf.append({"text":p.text.strip(),"align":str(p.alignment),"style":p.style.name,"runs":runs})
   cells.append(pf)
  rows.append(cells)
 tabs.append({"table":ti,"style":t.style.name if t.style else None,"rows":rows[:2]})
out={"headings":paras,"tables":tabs}
open(r"tmp/format_inspection.json","w",encoding="utf-8").write(json.dumps(out,ensure_ascii=False,indent=2))
