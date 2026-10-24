from docx import Document
from docx.oxml.ns import qn
import fitz,json,os
report=Document(r"deliverables/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO.docx")
breaks=[]
for i,p in enumerate(report.paragraphs):
 b=[x for x in p._p.iter(qn("w:br")) if x.get(qn("w:type"))=="page"]
 if b:breaks.append({"i":i,"text":p.text[:150]})
sections=[]
for i,e in enumerate(report.element.body):
 if e.tag==qn("w:sectPr"):sections.append({"bodyIndex":i,"kind":"sectPr"})
 elif e.tag==qn("w:p"):
  p=next((p for p in report.paragraphs if p._p is e),None)
  if p and p.text.strip().startswith(("PHẦN ","CHƯƠNG ")):
   sections.append({"bodyIndex":i,"text":p.text.strip(),"pStyle":p.style.name})
guide=r"C:/Users/dangv/Downloads/Hướng Dẫn Định Dạng Tài Liệu KLTN_ Sửa Style và Thiết Lập Trang - Studocu.pdf"
pages=[]
if os.path.exists(guide):
 pdf=fitz.open(guide)
 for i,p in enumerate(pdf):
  t=p.get_text()
  if any(k.lower() in t.lower() for k in ["bảng", "table", "font", "cỡ chữ"]):
   lines=[x.strip() for x in t.splitlines() if any(k.lower() in x.lower() for k in ["bảng","table","font","cỡ chữ","times new roman"])]
   if lines:pages.append({"page":i+1,"lines":lines[:30]})
sample_path=r"C:/Users/dangv/.gemini/antigravity/scratch/bao_cao_thuc_tap/backups/backup_20261023_125501/Bao_Cao_Thuc_Tap_Tot_Nghiep_Truong_Trung_Kien.docx"
sample=[]
if os.path.exists(sample_path):
 sd=Document(sample_path)
 for ti,t in enumerate(sd.tables[:4]):
  cells=[]
  for row in t.rows[:2]:
   cells.append([{"text":c.text[:90],"paras":[{"align":str(p.alignment),"runs":[{"font":r.font.name,"size":r.font.size.pt if r.font.size else None,"bold":r.bold,"italic":r.italic} for r in p.runs if r.text.strip()]} for p in c.paragraphs if p.text.strip()]} for c in row.cells[:3]])
  sample.append({"table":ti,"style":t.style.name if t.style else None,"rows":cells})
open(r"tmp/format_rules_research.json","w",encoding="utf-8").write(json.dumps({"breaks":breaks,"partsAndChapters":sections,"guideMatches":pages,"sampleTables":sample},ensure_ascii=False,indent=2))
