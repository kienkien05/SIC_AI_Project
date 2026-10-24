from docx import Document
import fitz,json,os
doc=Document(r"tmp/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO_final.docx")
pdf=fitz.open(r"tmp/Bao_Cao_Thuc_Tap_Tot_Nghiep_FPT_YOLO_final.pdf")
figs=[]
for p in doc.paragraphs:
 if p.text.strip().startswith(("Hình 1.1.","Hình 1.2.")):
  figs.append({"text":p.text.strip(),"italic":all(r.italic is True for r in p.runs if r.text.strip()),"centered":p.alignment==1})
all_text="\n".join(p.text for p in doc.paragraphs)
md=open(r"deliverables/FPT_Internship_Report_Content_Revised.md",encoding="utf-8").read()
chart_pages=[]
for i,p in enumerate(pdf):
 t=p.get_text()
 if "Hình 1.1." in t or "Hình 1.2." in t:chart_pages.append({"page":i+1,"captions":[x for x in t.splitlines() if x.startswith("Hình 1.")]})
report={"docx_pages":None,"pdf_pages":len(pdf),"inline_images":len(doc.inline_shapes),"figure_captions":figs,"reference_11":any(p.text.startswith("11. FPT Smart Cloud.") for p in doc.paragraphs),"old_defensive_phrases":{k:(k in all_text) for k in ["Mẫu báo cáo yêu cầu","không nêu tên cá nhân hoặc chi tiết cơ cấu chưa được cho phép","Tài liệu hội nhập không phải báo cáo tài chính"]},"md_image_links":md.count("!["),"md_chart_files_exist":[os.path.exists(r"deliverables/images/FPT_Smart_Cloud_Org_2026.png"),os.path.exists(r"deliverables/images/FPT_AI_Product_Org_2026.png")],"chart_pages":chart_pages,"toc_chart_rows":sum(1 for p in pdf if "Hình 1.1. Sơ đồ tổ chức" in p.get_text() or "Hình 1.2. Cơ cấu tổ chức" in p.get_text())}
open("tmp/final_validation.json","w",encoding="utf-8").write(json.dumps(report,ensure_ascii=False,indent=2))
