from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB

ROOT = Path(r'C:\Users\dangv\Downloads')
OUT = Path.cwd()

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9')

def add_table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E79'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2: shade(cells[i],'EAF2F8')
    borders(t); return t

def heading(doc, text, level=1):
    p=doc.add_paragraph(text, style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; return p

def para(doc, text, boldlead=None):
    p=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15; return p

def page(doc): doc.add_page_break()

def build_report():
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10.5)
    for n,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
        styles[n].font.name='Arial'; styles[n].font.size=Pt(size); styles[n].font.bold=True; styles[n].font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('SAMSUNG INNOVATION CAMPUS\n').bold=True; p.add_run('AI COURSE\n\n').bold=True
    r=p.add_run('FINAL REPORT\n'); r.bold=True; r.font.size=Pt(22)
    p.add_run('Smart Passive Attendance System\n\n').bold=True; p.add_run('SPAS - AI Based Face Attendance System')
    page(doc)
    heading(doc,'Executive Summary',1)
    para(doc,'SPAS is an AI-based attendance system that detects faces from a classroom camera, converts each detected face into a numerical embedding, compares the embedding with registered students, and records attendance evidence. The project focuses on a practical AI pipeline: data preparation, face detection, face recognition, evaluation, and integration into a web system.')
    para(doc,'The solution separates detection and recognition. YOLO11n locates faces. An ArcFace-style embedding model produces a 512-dimensional vector. Cosine similarity then matches the vector with a class gallery. The system returns a known identity or UNKNOWN when confidence is insufficient.')
    add_table(doc,['Item','Summary'],[['Problem','Manual attendance is slow and difficult to verify.'],['AI approach','YOLO11n detection plus 512D face embedding and cosine matching.'],['System output','Identity, confidence, attendance status, timestamp, and evidence image.'],['Current scope','MVP for classroom attendance; behavior monitoring and large-scale multi-camera support remain future work.']])
    page(doc)
    heading(doc,'Table of Contents',1)
    for x in ['Introduction','Chapter 1. Theoretical Background','Chapter 2. System Analysis and Design','Chapter 3. Model Training and System Integration','Chapter 4. Results Testing and Evaluation','Conclusion and Future Development','References','Appendices']:
        doc.add_paragraph(x)
    page(doc)
    heading(doc,'Introduction',1)
    heading(doc,'1. Problem Statement',2)
    para(doc,'Attendance is a recurring process in education. Manual roll calls consume class time, create transcription errors, and provide limited evidence when a result is disputed. A camera-based system can reduce repeated manual work while preserving a teacher review step.')
    heading(doc,'2. Objectives',2)
    para(doc,'The project objectives are to build a face attendance pipeline, evaluate the AI models on held-out data, integrate the models with a web application, and provide evidence that a teacher can review.')
    heading(doc,'3. Scope and Limitations',2)
    para(doc,'The MVP supports registered students, classroom webcam or camera input, one class roster, attendance sessions, unknown-person rejection, and evidence storage. It does not claim perfect recognition, continuous behavior analysis, or production-scale multi-camera performance.')
    heading(doc,'4. Report Structure',2)
    para(doc,'Chapter 1 introduces the theory. Chapter 2 describes requirements and architecture. Chapter 3 explains the datasets, preprocessing, model configuration, training, metrics, and integration. Chapter 4 presents measured results, functional tests, limitations, and evaluation.')
    page(doc)
    heading(doc,'Chapter 1. Theoretical Background',1)
    heading(doc,'1.1 Computer Vision',2); para(doc,'Computer vision converts images or video into structured information. In this project, the structured information is a face bounding box, an identity embedding, and an attendance event.')
    heading(doc,'1.2 Face Detection',2); para(doc,'Face detection is an object detection task. The detector predicts the location of each face and a confidence score. Detection happens before recognition because the recognition model should receive a cropped face rather than the whole classroom image.')
    heading(doc,'1.3 Convolutional Neural Networks',2); para(doc,'CNN layers learn spatial patterns through convolution filters. Early layers learn edges and textures. Deeper layers learn more complex face patterns. Pooling or strided convolution reduces spatial size and increases the receptive field.')
    heading(doc,'1.4 Transfer Learning',2); para(doc,'Transfer learning starts from weights learned on a large dataset, then adapts the model to a smaller domain dataset. This reduces training time and data requirements compared with training from random initialization.')
    heading(doc,'1.5 Face Embedding',2); para(doc,'An embedding is a compact numeric representation of a face. Images of the same person should have nearby vectors, while different identities should be separated. SPAS uses a 512-dimensional embedding.')
    heading(doc,'1.6 Cosine Similarity',2); para(doc,'Cosine similarity measures the angle between two vectors. After normalization, a higher score indicates greater similarity. The system compares a query embedding with each registered gallery embedding and rejects the match when the score is below the calibrated threshold.')
    page(doc)
    heading(doc,'Chapter 1. Theoretical Background Continued',1)
    heading(doc,'1.7 ArcFace Principle',2); para(doc,'ArcFace adds an angular margin to classification training. The margin forces samples from the same identity to form compact clusters and pushes different identities apart on a normalized hypersphere. In deployment, the classification head is not used for naming a new student; the embedding is matched with a gallery.')
    heading(doc,'1.8 Object Detection Metrics',2); para(doc,'Precision measures the fraction of predicted faces that are correct. Recall measures the fraction of ground-truth faces that are found. mAP50 uses an IoU threshold of 0.50. mAP50-95 averages performance across IoU thresholds from 0.50 to 0.95.')
    heading(doc,'1.9 Recognition Metrics',2); para(doc,'Top-1 accuracy records whether the highest-ranked identity is correct. Top-5 accuracy records whether the correct identity is among five candidates. False acceptance, false rejection, UNKNOWN rejection, cosine score, and score margin are important for attendance because an incorrect identity is more harmful than an unanswered sample.')
    heading(doc,'1.10 Ethical and Privacy Considerations',2); para(doc,'Face data is sensitive biometric data. The system should collect consent, limit access, protect stored images and embeddings, avoid logging raw face data, define retention periods, and preserve teacher correction. AI output is evidence for attendance, not an unquestionable decision.')
    page(doc)
    heading(doc,'Chapter 2. System Analysis and Design',1)
    heading(doc,'2.1 Stakeholders',2); add_table(doc,['Stakeholder','Need'],[['Student','Register face and view attendance history.'],['Teacher','Start attendance, review results, and correct exceptional cases.'],['Administrator','Manage users, classes, cameras, and biometric enrollment.'],['AI service','Detect faces, generate embeddings, and return matches.']])
    heading(doc,'2.2 Functional Requirements',2)
    for s in ['User authentication and role-based access','Student enrollment with multiple face frames','Class roster management','Teacher attendance session','Face detection and identity matching','UNKNOWN_PERSON handling','Evidence image and attendance history','Teacher correction and review'] : doc.add_paragraph(s,style='List Bullet')
    heading(doc,'2.3 Non-functional Requirements',2); para(doc,'The system should provide clear error messages, protect service-to-service requests, keep the database inaccessible from the public network, and remain usable when a face cannot be confidently recognized.')
    page(doc)
    heading(doc,'Chapter 2. System Analysis and Design Continued',1)
    heading(doc,'2.4 Architecture',2); para(doc,'The Frontend sends authenticated requests to the Backend. The Backend applies authorization and attendance business rules, then calls the AI service through an internal HTTP request. The AI service runs detection and recognition. The Backend stores attendance and evidence, and the Frontend displays the result.')
    add_table(doc,['Layer','Responsibility'],[['Frontend','User interface, webcam capture, attendance display.'],['Backend','Authentication, authorization, attendance rules, persistence.'],['AI service','YOLO11n detection, embedding generation, matching.'],['Database','Users, classes, rosters, sessions, logs, evidence metadata.']])
    heading(doc,'2.5 Attendance Flow',2); para(doc,'The teacher selects a class and starts a session. The client captures frames during the scan window. Each frame is sent to the Backend. The Backend calls the AI service. Recognized identities are merged across frames, duplicates are suppressed, and the final session is saved with evidence.')
    heading(doc,'2.6 Unknown and Duplicate Handling',2); para(doc,'The system rejects low-similarity matches as UNKNOWN. It also prevents one face enrollment from being assigned to multiple accounts. These controls reduce false attendance and support administrative review.')
    page(doc)
    heading(doc,'Chapter 2. System Analysis and Design Continued',1)
    heading(doc,'2.7 Data Model',2); para(doc,'Core entities are User, Class, Enrollment, FaceEmbedding, AttendanceSession, AttendanceRecord, and Evidence. Enrollment belongs to a user. A class contains a roster. A session belongs to a class and produces attendance records. Evidence stores a reference to the captured image and not unnecessary duplicate biometric data.')
    heading(doc,'2.8 API Integration',2); para(doc,'The Frontend uses the versioned Backend API. The Backend sends a protected request to the AI service and validates the response before writing attendance. The Frontend does not call the model directly. This design keeps model credentials and business rules on the server side.')
    heading(doc,'2.9 Design Decisions',2); add_table(doc,['Decision','Reason'],[['Separate detector and recognizer','Allows independent evaluation and replacement.'],['Embedding gallery','New students can enroll without retraining a name classifier.'],['Multi-frame scan','Reduces sensitivity to one blurred or occluded frame.'],['Teacher review','Handles uncertainty and exceptional classroom conditions.']])
    page(doc)
    heading(doc,'Chapter 3. Model Training and System Integration',1)
    heading(doc,'3.1 AI Pipeline Overview',2); para(doc,'The AI workflow is organized into EDA, preprocessing, training, validation, metric export, real-image inference, model packaging, and service integration. The two supplied notebooks implement this workflow in runnable cells.')
    heading(doc,'3.2 Dataset Preparation',2); para(doc,'Images are enumerated, invalid files are excluded, identities are counted, and examples are displayed before training. The split must be identity- or session-aware where possible. Images from the same capture burst should not appear in both training and validation sets because that inflates the reported score.')
    heading(doc,'3.3 YOLO11n Detector',2); para(doc,'The detector notebook uses YOLO11n with image size 640, 25 epochs, and batch size 8. The training output includes validation metrics and a real-image inference output. The detector returns bounding boxes and confidence scores for faces.')
    heading(doc,'3.4 ArcFace Recognition Model',2); para(doc,'The recognition notebook resizes input faces to 160 × 160. InceptionResnetV1 produces a 512-dimensional embedding. An ArcFace head is used during training with normalized embeddings and angular separation. The deployment path uses the embedding model and cosine matching against the gallery.')
    page(doc)
    heading(doc,'Chapter 3. Model Training and System Integration Continued',1)
    heading(doc,'3.5 Training Configuration',2); add_table(doc,['Parameter','YOLO11n','ArcFace recognition'],[['Input size','640','160 × 160'],['Epochs','25','25'],['Batch size','8','64'],['Output','Boxes and confidence','512D embedding'],['Validation','Precision, Recall, mAP','Top-1, Top-5, cosine metrics']])
    heading(doc,'3.6 Preprocessing',2); para(doc,'Detection preprocessing preserves the image format expected by YOLO. Recognition preprocessing converts the crop to RGB, resizes it to 160 × 160, and applies the same normalization used during training. Keeping training and serving preprocessing identical is required for valid similarity scores.')
    heading(doc,'3.7 Loss and Optimization',2); para(doc,'The recognition notebook uses cross-entropy over the ArcFace head with label smoothing and AdamW optimization. A cosine annealing scheduler changes the learning rate across epochs. The best checkpoint is selected using validation performance rather than the final epoch alone.')
    heading(doc,'3.8 Model Packaging',2); para(doc,'The selected weights, preprocessing parameters, model version, threshold, and gallery format must be packaged together. A model file without its preprocessing and threshold is not a reproducible deployment artifact.')
    page(doc)
    heading(doc,'Chapter 3. Model Training and System Integration Continued',1)
    heading(doc,'3.9 Evaluation Protocol',2); para(doc,'Evaluation uses images not used for training. For detection, predictions are compared with labeled bounding boxes. For recognition, the gallery is built from enrollment images and queries are evaluated against known and unknown identities. Results are reported overall and by scenario such as frontal, pose, mask, low light, and blur.')
    heading(doc,'3.10 Metric Status',2); para(doc,'The supplied notebooks contain code that writes metrics JSON files, but final numeric values must be inserted only after the fine-tuning run finishes. The report must not present training accuracy as production recognition accuracy.')
    add_table(doc,['Metric group','Values to report'],[['YOLO11n','Precision, Recall, mAP50, mAP50-95'],['Recognition','Best validation Top-1, Top-5, gallery Top-1, gallery Top-5'],['Matching','Mean cosine similarity, minimum cosine, top-1/top-2 margin'],['Runtime','Average inference time per image/frame']])
    heading(doc,'3.11 Integration with SPAS',2); para(doc,'The AI service loads the detector and recognizer at startup. The Backend sends a frame or face observation, receives structured predictions, checks the service response, applies roster matching and attendance rules, and persists the result. The Frontend renders bounding boxes, identity, score, and evidence.')
    page(doc)
    heading(doc,'Chapter 4. Results Testing and Evaluation',1)
    heading(doc,'4.1 Experimental Environment',2); add_table(doc,['Component','Configuration'],[['Training','Kaggle GPU notebook; final run configuration recorded in notebook output.'],['Serving','Dockerized SPAS services.'],['Input','Webcam or classroom camera frame.'],['Output','Attendance result, score, and evidence.']])
    heading(doc,'4.2 Detection Results',2); para(doc,'Insert the final YOLO11n validation table after the run completes. Include the dataset split, number of images, number of labeled faces, confidence threshold, and the four detection metrics. Add at least one figure showing a correct detection and one showing a difficult case.')
    add_table(doc,['Metric','Final value','Interpretation'],[['Precision','[pending run]','False positive control.'],['Recall','[pending run]','Missed-face control.'],['mAP50','[pending run]','Detection quality at IoU 0.50.'],['mAP50-95','[pending run]','Stricter localization quality.']])
    heading(doc,'4.3 Recognition Results',2); add_table(doc,['Metric','Final value','Interpretation'],[['Validation Top-1','[pending run]','Highest-ranked identity accuracy.'],['Validation Top-5','[pending run]','Correct identity within five candidates.'],['Gallery Top-1','[pending run]','Direct gallery matching result.'],['Mean cosine similarity','[pending run]','Similarity strength across queries.']])
    page(doc)
    heading(doc,'Chapter 4. Results Testing and Evaluation Continued',1)
    heading(doc,'4.4 Functional Test Plan',2); add_table(doc,['ID','Scenario','Expected result'],[['T01','Valid login','User reaches role-appropriate page.'],['T02','Enroll valid face','Enrollment is saved with embedding.'],['T03','Start class attendance','Camera scan begins.'],['T04','Known student','Identity and attendance are recorded.'],['T05','Unknown person','System returns UNKNOWN and does not mark roster.'],['T06','Duplicate enrollment','System rejects the face and explains the conflict.'],['T07','Evidence review','Teacher can inspect captured evidence.'],['T08','Teacher correction','Updated status is persisted and visible in history.']])
    heading(doc,'4.5 System Evaluation',2); para(doc,'Evaluate the complete system with a fixed test script. Record total frames, detected faces, known matches, unknown matches, false matches, rejected frames, average response time, and failures. Report both successful behavior and failure cases.')
    heading(doc,'4.6 Error Analysis',2); para(doc,'Typical errors include blur, poor lighting, extreme pose, occlusion, small faces, similar-looking people, and stale gallery data. For each error, show an example, identify whether detection or recognition caused it, and state the mitigation.')
    page(doc)
    heading(doc,'Chapter 4. Results Testing and Evaluation Continued',1)
    heading(doc,'4.7 Comparison with Manual Attendance',2); para(doc,'Compare the time required to complete one attendance session, the number of manual corrections, and the availability of evidence. The comparison should use the same class size and session conditions. Do not claim general speed improvement from a single demonstration.')
    heading(doc,'4.8 Limitations',2); para(doc,'The MVP depends on camera placement, enrollment quality, lighting, and the representativeness of the dataset. A score threshold calibrated on one camera may not transfer to another. The model does not guarantee liveness detection against replay attacks unless a dedicated liveness model is added.')
    heading(doc,'4.9 Security and Privacy Evaluation',2); para(doc,'Check that protected endpoints reject missing or invalid tokens, the AI service requires its service key, PostgreSQL is not publicly exposed, and logs do not contain raw biometric vectors. Confirm that teachers can review and correct AI output.')
    heading(doc,'4.10 Discussion',2); para(doc,'The system is successful when it provides a reproducible AI pipeline and a usable attendance workflow. The strongest evidence combines model metrics, scenario-based inference, functional tests, and clear limitations. A single high accuracy number is not sufficient evidence for deployment.')
    page(doc)
    heading(doc,'Conclusion and Future Development',1)
    para(doc,'This project developed an end-to-end AI attendance prototype. The system combines YOLO11n face detection, a 512-dimensional face embedding model, cosine similarity, roster matching, attendance persistence, and evidence review. The design keeps the AI service separate from business logic so that the models can be evaluated and replaced independently.')
    heading(doc,'Future Development',2)
    for s in ['Complete metric runs and calibrate thresholds on camera-specific validation data.','Add dedicated liveness or anti-spoofing evaluation.','Expand testing across lighting, pose, occlusion, and unknown identities.','Add temporal tracking and efficient multi-face batching.','Add privacy controls, retention policy, consent records, and audit review.','Evaluate multi-camera deployment and GPU or edge inference only after profiling.'] : doc.add_paragraph(s,style='List Bullet')
    heading(doc,'References',1)
    for s in ['Ultralytics. YOLO Documentation and YOLO11 model documentation.','Deng, J. et al. ArcFace: Additive Angular Margin Loss for Deep Face Recognition.','Schroff, F. et al. FaceNet: A Unified Embedding for Face Recognition and Clustering.','Samsung Innovation Campus. AI Course Capstone Project Final Report Template.','Supplied SPAS project notebooks: YOLO11n detector and ArcFace recognition.'] : doc.add_paragraph(s)
    heading(doc,'Appendix A. Metric Replacement Checklist',1)
    for s in ['Replace every [pending run] value with a value from the final saved JSON.','Record dataset version and split.','Record model checkpoint and threshold.','Attach confusion examples and real-image inference outputs.','Run the complete functional test table and mark Pass or Fail.'] : doc.add_paragraph(s,style='List Bullet')
    doc.save(OUT/'SPAS_Final_Report.docx')

def build_slides():
    prs=Presentation(ROOT/'SIC_AI_Capstone Project_Presentation Slide Template.pptx')
    # retain template cover, replace visible text in all existing shapes
    texts=[('Team Name','SIC AI Team'),('Project\nName','Smart Passive Attendance System'),('Project Name','Smart Passive Attendance System'),('Slide Title','AI Attendance Pipeline')]
    for slide, vals in zip(prs.slides,texts):
        for sh in slide.shapes:
            if hasattr(sh,'text') and sh.text.strip():
                sh.text=vals[1]
    def add(title, body, accent='1F4E79'):
        s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=PRGB(248,250,252)
        bar=s.shapes.add_shape(1,0,0,PInches(.22),prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb=PRGB.from_string(accent); bar.line.fill.background()
        tx=s.shapes.add_textbox(PInches(.7),PInches(.45),PInches(12),PInches(.65)); tf=tx.text_frame; tf.text=title; tf.paragraphs[0].font.size=PPt(28); tf.paragraphs[0].font.bold=True; tf.paragraphs[0].font.color.rgb=PRGB(15,23,42)
        box=s.shapes.add_textbox(PInches(.9),PInches(1.45),PInches(11.6),PInches(5.5)); tf=box.text_frame; tf.word_wrap=True
        for i,line in enumerate(body.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.level=0; p.font.size=PPt(20 if len(body.split('\n'))<7 else 17); p.font.color.rgb=PRGB(30,41,59); p.space_after=PPt(12)
        return s
    add('Why SPAS','Manual roll call consumes class time\nFace attendance provides repeatable evidence\nGoal: practical AI pipeline from camera to attendance record')
    add('System Architecture','Webcam / camera\nYOLO11n face detection\n512D face embedding\nCosine similarity with class gallery\nBackend saves attendance and evidence')
    add('Model 1: YOLO11n Detector','Input size: 640\nTraining: 25 epochs, batch size 8\nOutput: face bounding boxes and confidence\nMetrics: Precision, Recall, mAP50, mAP50-95')
    add('Model 2: ArcFace Recognition','Input: 160 × 160 face crop\nInceptionResnetV1 embedding: 512 dimensions\nArcFace margin during training\nGallery matching with cosine similarity')
    add('Training Pipeline','EDA → preprocessing → train → validate\nSave best checkpoint and metrics JSON\nRun real-image inference\nPackage weights, preprocessing, threshold, and gallery')
    add('Integration Flow','Frontend starts attendance session\nBackend authorizes request and calls AI service\nAI returns boxes, identity, score\nBackend deduplicates and stores result\nTeacher reviews evidence')
    add('Evaluation Plan','Detection: Precision, Recall, mAP50, mAP50-95\nRecognition: Top-1, Top-5, cosine score, margin\nSystem: known, unknown, false match, response time\nFinal values inserted after fine-tuning completes')
    add('Functional Tests','Login and role access\nFace enrollment\nKnown student attendance\nUNKNOWN_PERSON rejection\nDuplicate enrollment rejection\nEvidence and teacher correction')
    add('Limitations and Future Work','Sensitive to lighting, pose, blur, and camera placement\nThreshold needs camera-specific calibration\nLiveness and anti-spoofing need dedicated evaluation\nFuture: tracking, batching, multi-camera, privacy controls')
    add('Conclusion','SPAS demonstrates an end-to-end AI attendance workflow\nDetection and recognition are independently evaluable\nSystem output remains reviewable by teachers\nNext milestone: complete final metrics and scenario report')
    prs.save(OUT/'SPAS_Presentation.pptx')

if __name__=='__main__': build_report(); build_slides()
